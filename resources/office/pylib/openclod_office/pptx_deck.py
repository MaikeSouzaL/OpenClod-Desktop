# -*- coding: utf-8 -*-
"""pptx_deck - a clean 16:9 presentation with native, editable charts and tables (python-pptx underneath).

    from openclod_office.pptx_deck import Deck
    deck = Deck(footer="Preparado pela Equipe Alfa", accent="1F4E79")
    deck.title_slide("Resultados de vendas 2026", "Diretoria Comercial")
    deck.kpi_slide("Números do ano", [("Receita", "R$ 639.377"), ("Meta", "133%"), ("Média mensal", "R$ 106.563")])
    deck.chart_slide("Receita por região", ["Sul", "Sudeste"], {"Receita": [150000, 260000]}, kind="column", number_format='"R$" #,##0',
                     insight=["**Sudeste** lidera", "Sul bate a meta"])
    deck.bullets_slide("Próximos passos", ["Rever o preço do Plano Básico", "Abrir canal online"], notes="Fala do apresentador")
    deck.save(r"C:\\pasta\\apresentacao.pptx")

Every slide keeps a real title placeholder (outline, navigation, accessibility). `deck.prs` is the python-pptx Presentation and
`deck.blank_slide(title)` returns a titled slide to draw on by hand. Text accepts **bold** inline.
"""
import math
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

from .common import DEFAULT_ACCENT, RIGHT_ALIGNED, fmt_value, hex_to_rgb, lighten, split_bold

SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN = 0.65
CONTENT_TOP = 1.65
CONTENT_BOTTOM = 6.75
PALETTE = ["1F4E79", "2E75B6", "9DC3E6", "ED7D31", "A5A5A5", "70AD47"]
CHARTS = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}


def _rgb(color):
    return RGBColor(*hex_to_rgb(color))


def _lines_needed(text, width_in, size):
    per_line = max(1.0, (width_in * 72.0 - 14.0) / (size * 0.52))
    return max(1, math.ceil(len(text) / per_line)) if text else 1


def _fit(texts, width_in, height_in, sizes=(28, 24, 22, 20, 18, 16, 14), gap_pt=8):
    """Largest font size from `sizes` at which the paragraphs fit the box (estimate); the smallest when none does."""
    for size in sizes:
        total = sum(_lines_needed(t, width_in, size) * size * 1.2 + gap_pt for t in texts)
        if total <= (height_in * 72.0 - 7.0):
            return size
    return sizes[-1]


def _runs(paragraph, text, size, color="262626", bold=False, font=None):
    for piece, strong in split_bold(text):
        run = paragraph.add_run()
        run.text = piece
        run.font.size = Pt(size)
        run.font.bold = bool(strong or bold)
        run.font.color.rgb = _rgb(color)
        if font:
            run.font.name = font


def _bullet(paragraph, level, size, accent):
    """Real bullet formatting for a text box paragraph (hanging indent, accent-colored dot)."""
    ppr = paragraph._p.get_or_add_pPr()
    indent = int(size * 12700 * 1.1)
    ppr.set("marL", str(indent * (level + 1)))
    ppr.set("indent", str(-indent))
    for tag in ("a:buNone", "a:buChar", "a:buClr", "a:buFont"):
        for old in ppr.findall(qn(tag)):
            ppr.remove(old)
    color = etree.SubElement(ppr, qn("a:buClr"))
    etree.SubElement(color, qn("a:srgbClr")).set("val", accent)
    font = etree.SubElement(ppr, qn("a:buFont"))
    font.set("typeface", "Arial")
    char = etree.SubElement(ppr, qn("a:buChar"))
    char.set("char", "•" if level == 0 else "–")


class Deck:
    def __init__(self, title=None, footer=None, accent=DEFAULT_ACCENT, font="Calibri"):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
        self.accent = accent
        self.font = font
        self.footer = footer
        self._covers = set()
        if title:
            self.prs.core_properties.title = title

    # ------------------------------------------------------------------ building blocks

    def _notes(self, slide, notes):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def _title(self, slide, text, size=30, color="262626"):
        shape = slide.shapes.title
        shape.left, shape.top, shape.width, shape.height = Inches(MARGIN), Inches(0.42), Inches(SLIDE_W - 2 * MARGIN), Inches(1.0)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        size = _fit([text], SLIDE_W - 2 * MARGIN, 1.0, sizes=(size, size - 4, size - 8), gap_pt=0)
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        _runs(paragraph, text, size, color, bold=True, font=self.font)
        return shape

    def _band(self, slide):
        band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, self.prs.slide_width, Inches(0.14))
        band.fill.solid()
        band.fill.fore_color.rgb = _rgb(self.accent)
        band.line.fill.background()
        band.name = "Faixa"
        return band

    def blank_slide(self, title=None, notes=None):
        """A slide with the accent band and (when given) a real title; draw the rest with python-pptx (`slide.shapes`)."""
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[5])
        self._band(slide)
        if title:
            self._title(slide, title)
        else:
            slide.shapes.title._element.getparent().remove(slide.shapes.title._element)
        self._notes(slide, notes)
        return slide

    def _textbox(self, slide, x, y, w, h, name=None):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        box.text_frame.word_wrap = True
        box.text_frame.auto_size = MSO_AUTO_SIZE.NONE
        if name:
            box.name = name
        return box

    def _bullets_into(self, box, items, width_in, height_in, sizes=(26, 24, 22, 20, 18, 16, 14)):
        flat = []
        for item in items:
            text, subs = (item if isinstance(item, tuple) else (item, []))
            flat.append(text)
            flat.extend(subs)
        size = _fit(flat, width_in - 0.4, height_in, sizes=sizes)
        frame = box.text_frame
        first = True
        for item in items:
            text, subs = (item if isinstance(item, tuple) else (item, []))
            paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            _runs(paragraph, text, size, font=self.font)
            _bullet(paragraph, 0, size, self.accent)
            paragraph.space_after = Pt(8)
            for sub in subs:
                sp = frame.add_paragraph()
                _runs(sp, sub, max(size - 4, 12), "595959", font=self.font)
                _bullet(sp, 1, max(size - 4, 12), self.accent)
                sp.space_after = Pt(4)

    # ------------------------------------------------------------------ slide types

    def title_slide(self, title, subtitle=None, notes=None):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[0])
        self._covers.add(len(self.prs.slides) - 1)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(self.accent)
        shape = slide.shapes.title
        shape.left, shape.top, shape.width, shape.height = Inches(0.9), Inches(2.1), Inches(SLIDE_W - 1.8), Inches(1.8)
        frame = shape.text_frame
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.BOTTOM
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        _runs(paragraph, title, _fit([title], SLIDE_W - 1.8, 1.8, sizes=(46, 40, 34, 30), gap_pt=0), "FFFFFF", bold=True, font=self.font)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(4.08), Inches(1.4), Inches(0.07))
        rule.fill.solid()
        rule.fill.fore_color.rgb = _rgb(lighten(self.accent, 0.7))
        rule.line.fill.background()
        sub = slide.placeholders[1]
        sub.left, sub.top, sub.width, sub.height = Inches(0.9), Inches(4.35), Inches(SLIDE_W - 1.8), Inches(1.3)
        sub.text_frame.word_wrap = True
        sub.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        sp = sub.text_frame.paragraphs[0]
        sp.alignment = PP_ALIGN.LEFT
        if subtitle:
            _runs(sp, subtitle, 22, lighten(self.accent, 0.8), font=self.font)
        else:
            sub._element.getparent().remove(sub._element)
        self._notes(slide, notes)
        return slide

    def section_slide(self, title, subtitle=None, notes=None):
        return self.title_slide(title, subtitle, notes)

    def bullets_slide(self, title, bullets, notes=None, lead=None):
        """bullets: list of 'text' or ('text', ['sub', 'sub']). Up to ~6 short bullets; `lead` is a one-line takeaway above them."""
        slide = self.blank_slide(title, notes)
        top = CONTENT_TOP
        if lead:
            box = self._textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, 0.7, "Destaque")
            _runs(box.text_frame.paragraphs[0], lead, 20, self.accent, bold=True, font=self.font)
            top += 0.85
        box = self._textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - top, "Tópicos")
        self._bullets_into(box, bullets, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - top)
        return slide

    def kpi_slide(self, title, kpis, notes=None, bullets=None):
        """kpis: up to 4 of (label, value) or (label, value, detail). `bullets` (optional) explains them below the cards."""
        slide = self.blank_slide(title, notes)
        count = len(kpis)
        gap = 0.3
        width = (SLIDE_W - 2 * MARGIN - gap * (count - 1)) / count
        height = 2.05
        for index, kpi in enumerate(kpis):
            label, value = kpi[0], kpi[1]
            detail = kpi[2] if len(kpi) > 2 else None
            x = MARGIN + index * (width + gap)
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(CONTENT_TOP), Inches(width), Inches(height))
            card.adjustments[0] = 0.06
            card.fill.solid()
            card.fill.fore_color.rgb = _rgb(lighten(self.accent, 0.92))
            card.line.color.rgb = _rgb(lighten(self.accent, 0.75))
            card.name = "Indicador %d" % (index + 1)
            frame = card.text_frame
            frame.word_wrap = True
            frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            vsize = _fit([str(value)], width - 0.3, 0.9, sizes=(40, 34, 30, 26, 22), gap_pt=0)
            p1 = frame.paragraphs[0]
            p1.alignment = PP_ALIGN.CENTER
            _runs(p1, str(value), vsize, self.accent, bold=True, font=self.font)
            p2 = frame.add_paragraph()
            p2.alignment = PP_ALIGN.CENTER
            _runs(p2, str(label), 15, "404040", font=self.font)
            if detail:
                p3 = frame.add_paragraph()
                p3.alignment = PP_ALIGN.CENTER
                _runs(p3, str(detail), 12, "7F7F7F", font=self.font)
        if bullets:
            top = CONTENT_TOP + height + 0.4
            box = self._textbox(slide, MARGIN, top, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - top, "Leitura")
            self._bullets_into(box, bullets, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - top, sizes=(22, 20, 18, 16))
        return slide

    def chart_slide(self, title, categories, series, kind="column", notes=None, number_format="#,##0", insight=None, colors=None):
        """
        categories: x labels. series: {'Nome': [valores]} (or a list of (name, values)). kind: column, bar, line, stacked, pie, doughnut.
        insight: bullets (or one string) shown beside the chart; without it the chart uses the full width.
        """
        slide = self.blank_slide(title, notes)
        items = list(series.items()) if isinstance(series, dict) else list(series)
        data = CategoryChartData()
        data.categories = [str(c) for c in categories]
        for name, values in items:
            data.add_series(str(name), [float(v) for v in values])
        has_insight = bool(insight)
        chart_w = (SLIDE_W - 2 * MARGIN) * (0.62 if has_insight else 1.0)
        frame = slide.shapes.add_chart(CHARTS[kind], Inches(MARGIN), Inches(CONTENT_TOP), Inches(chart_w), Inches(CONTENT_BOTTOM - CONTENT_TOP), data)
        frame.name = "Gráfico"
        chart = frame.chart
        chart.font.size = Pt(12)
        chart.font.name = self.font
        chart.has_title = False
        multi = len(items) > 1 or kind in ("pie", "doughnut")
        chart.has_legend = multi
        if multi:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(12)
        palette = colors or PALETTE
        plot = chart.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.number_format = number_format
        labels.number_format_is_linked = False
        labels.font.size = Pt(12)
        if kind in ("column", "bar"):
            labels.position = XL_LABEL_POSITION.OUTSIDE_END
            plot.gap_width = 70
        elif kind == "line":
            labels.position = XL_LABEL_POSITION.ABOVE
        if kind in ("pie", "doughnut"):
            for index in range(len(categories)):
                point = plot.series[0].points[index]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = _rgb(palette[index % len(palette)])
        else:
            for index, series_obj in enumerate(plot.series):
                color = _rgb(palette[index % len(palette)])
                if kind == "line":
                    series_obj.format.line.color.rgb = color
                    series_obj.format.line.width = Pt(2.5)
                    series_obj.smooth = False
                else:
                    series_obj.format.fill.solid()
                    series_obj.format.fill.fore_color.rgb = color
            value_axis = chart.value_axis
            if kind in ("column", "stacked", "bar"):
                value_axis.minimum_scale = 0  # bars start at zero, or their heights lie (the viewer's auto scale may start at 200.000)
            value_axis.has_major_gridlines = True
            value_axis.major_gridlines.format.line.color.rgb = _rgb("D9D9D9")
            value_axis.format.line.fill.background()
            value_axis.tick_labels.font.size = Pt(11)
            value_axis.tick_labels.number_format = number_format
            value_axis.tick_labels.number_format_is_linked = False
            chart.category_axis.tick_labels.font.size = Pt(12)
            chart.category_axis.format.line.color.rgb = _rgb("BFBFBF")
        if has_insight:
            points = [insight] if isinstance(insight, str) else list(insight)
            x = MARGIN + chart_w + 0.3
            box = self._textbox(slide, x, CONTENT_TOP + 0.2, SLIDE_W - MARGIN - x, CONTENT_BOTTOM - CONTENT_TOP - 0.4, "Leitura")
            self._bullets_into(box, points, SLIDE_W - MARGIN - x, CONTENT_BOTTOM - CONTENT_TOP - 0.4, sizes=(22, 20, 18, 16, 14))
        return slide

    def table_slide(self, title, header, rows, formats=None, notes=None, col_widths=None, total=False):
        """A native table. formats: per column text/int/num/brl/brl0/pct/pct0 (numbers right-aligned). total=True: last row is a total."""
        slide = self.blank_slide(title, notes)
        cols = len(header)
        formats = list(formats or ["text"] * cols)
        n_rows = len(rows) + 1
        row_h = min(0.62, (CONTENT_BOTTOM - CONTENT_TOP) / n_rows)
        width = SLIDE_W - 2 * MARGIN
        frame = slide.shapes.add_table(n_rows, cols, Inches(MARGIN), Inches(CONTENT_TOP), Inches(width), Inches(row_h * n_rows))
        frame.name = "Tabela"
        table = frame.table
        table.horz_banding = False
        widths = col_widths or [width / cols] * cols
        for index, w in enumerate(widths):
            table.columns[index].width = Inches(w)
        font_size = 16 if n_rows <= 7 else 14 if n_rows <= 10 else 12
        for r in range(n_rows):
            table.rows[r].height = Inches(row_h)
            for c in range(cols):
                cell = table.cell(r, c)
                raw = header[c] if r == 0 else rows[r - 1][c]
                text = str(raw) if r == 0 else str(fmt_value(raw, formats[c]))
                is_total = total and r == n_rows - 1
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.margin_left = cell.margin_right = Inches(0.12)
                paragraph = cell.text_frame.paragraphs[0]
                paragraph.alignment = PP_ALIGN.RIGHT if formats[c] in RIGHT_ALIGNED else PP_ALIGN.LEFT
                cell.fill.solid()
                if r == 0:
                    cell.fill.fore_color.rgb = _rgb(self.accent)
                    _runs(paragraph, text, font_size, "FFFFFF", bold=True, font=self.font)
                elif is_total:
                    cell.fill.fore_color.rgb = _rgb(lighten(self.accent, 0.82))
                    _runs(paragraph, text, font_size, "262626", bold=True, font=self.font)
                else:
                    cell.fill.fore_color.rgb = _rgb(lighten(self.accent, 0.94) if r % 2 == 0 else "FFFFFF")
                    _runs(paragraph, text, font_size, "262626", font=self.font)
        return slide

    def image_slide(self, title, image_path, caption=None, notes=None):
        if not os.path.isfile(image_path):
            raise FileNotFoundError(image_path)
        slide = self.blank_slide(title, notes)
        max_w, max_h = SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP - (0.5 if caption else 0)
        picture = slide.shapes.add_picture(image_path, Inches(MARGIN), Inches(CONTENT_TOP))
        ratio = min(Emu(Inches(max_w)) / picture.width, Emu(Inches(max_h)) / picture.height, 1.0 if picture.width < Inches(max_w) else 9e9)
        picture.width, picture.height = int(picture.width * ratio), int(picture.height * ratio)
        picture.left = int((self.prs.slide_width - picture.width) / 2)
        if caption:
            box = self._textbox(slide, MARGIN, CONTENT_BOTTOM - 0.4, SLIDE_W - 2 * MARGIN, 0.4, "Legenda")
            paragraph = box.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.CENTER
            _runs(paragraph, caption, 14, "595959", font=self.font)
        return slide

    def two_columns_slide(self, title, left_title, left_bullets, right_title, right_bullets, notes=None):
        slide = self.blank_slide(title, notes)
        gap = 0.5
        width = (SLIDE_W - 2 * MARGIN - gap) / 2
        for index, (head, bullets) in enumerate(((left_title, left_bullets), (right_title, right_bullets))):
            x = MARGIN + index * (width + gap)
            heading = self._textbox(slide, x, CONTENT_TOP, width, 0.6, "Título da coluna %d" % (index + 1))
            _runs(heading.text_frame.paragraphs[0], head, 22, self.accent, bold=True, font=self.font)
            box = self._textbox(slide, x, CONTENT_TOP + 0.75, width, CONTENT_BOTTOM - CONTENT_TOP - 0.75, "Coluna %d" % (index + 1))
            self._bullets_into(box, bullets, width, CONTENT_BOTTOM - CONTENT_TOP - 0.75, sizes=(22, 20, 18, 16, 14))
        return slide

    # ------------------------------------------------------------------ output

    def save(self, path):
        total = len(self.prs.slides)
        for index, slide in enumerate(self.prs.slides):
            if index in self._covers:
                continue
            if self.footer:
                box = self._textbox(slide, MARGIN, 6.98, 8.0, 0.32, "Rodapé")
                _runs(box.text_frame.paragraphs[0], self.footer, 11, "7F7F7F", font=self.font)
            box = self._textbox(slide, SLIDE_W - MARGIN - 1.6, 6.98, 1.6, 0.32, "Número do slide")
            paragraph = box.text_frame.paragraphs[0]
            paragraph.alignment = PP_ALIGN.RIGHT
            _runs(paragraph, "%d / %d" % (index + 1, total), 11, "7F7F7F", font=self.font)
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        return path
