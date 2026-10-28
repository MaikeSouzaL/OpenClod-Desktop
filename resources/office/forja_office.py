#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""openclod-office - checks the document environment of Openclod and inspects what was delivered.

Runs on the embedded Python runtime (python-docx, openpyxl, python-pptx, XlsxWriter, Pillow, lxml), through the
`openclod-office` command that run_command puts on PATH:

  openclod-office check   [--json]
  openclod-office inspect <file.docx|.xlsx|.pptx> [--json] [--full]
  openclod-office pdf     <file> [--out <folder>]      converts with LibreOffice (when installed)
  openclod-office recalc  <file.xlsx>                  recalculates with LibreOffice and lists formula errors

`inspect` answers "what is really inside this file?" without a screenshot: headings, tables, formulas, charts,
slides, and the usual defects (empty placeholders, text that may overflow, numbers stored as text, formulas
without a cached value or with an error). It never modifies the file.
"""
import argparse
import importlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from importlib import metadata as md

LIBS = (
    ("python-docx", "docx"),
    ("openpyxl", "openpyxl"),
    ("python-pptx", "pptx"),
    ("XlsxWriter", "xlsxwriter"),
    ("Pillow", "PIL"),
    ("lxml", "lxml.etree"),
)
ERROR_VALUES = ("#DIV/0!", "#REF!", "#NAME?", "#VALUE!", "#N/A", "#NUM!", "#NULL!")
PLACEHOLDER_TEXT = re.compile(r"lorem ipsum|\bTODO\b|\bXXX\b|\[inserir|\[preencher|clique para adicionar|click to add", re.I)
MARKDOWN_LEAK = re.compile(r"(^|\n)\s*#{1,6} \S|\*\*[^*\n]+\*\*|(^|\n)\s*[-*] \[[ x]\]")
EMU_PER_INCH = 914400.0
EMU_PER_CM = 360000.0


def out(text=""):
    sys.stdout.write(text + "\n")


# --------------------------------------------------------------------------------------- environment


def find_soffice():
    candidates = [os.environ.get("OPENCLOD_SOFFICE"), shutil.which("soffice"), shutil.which("soffice.exe")]
    for var in ("ProgramFiles", "ProgramFiles(x86)", "LOCALAPPDATA"):
        base = os.environ.get(var)
        if base:
            candidates.append(os.path.join(base, "LibreOffice", "program", "soffice.exe"))
            candidates.append(os.path.join(base, "Programs", "LibreOffice", "program", "soffice.exe"))
    for item in candidates:
        if item and os.path.isfile(item):
            return item
    return None


def library_report():
    libs = {}
    for dist, module in LIBS:
        try:
            importlib.import_module(module)
            libs[dist] = md.version(dist)
        except Exception as exc:  # a missing or broken library is reported, not raised
            libs[dist] = None
            libs["_error_" + dist] = "%s: %s" % (type(exc).__name__, exc)
    return libs


def cmd_check(args):
    libs = library_report()
    missing = [dist for dist, _ in LIBS if not libs.get(dist)]
    soffice = find_soffice()
    report = {
        "python": sys.executable,
        "pythonVersion": sys.version.split()[0],
        "libraries": {dist: libs.get(dist) for dist, _ in LIBS},
        "missing": missing,
        "libreoffice": soffice,
        "ok": not missing,
    }
    if args.json:
        out(json.dumps(report, ensure_ascii=False, indent=2))
        return 0 if report["ok"] else 1
    out("Ambiente de documentos do Openclod")
    out("  Python embutido : %s (%s)" % (report["python"], report["pythonVersion"]))
    out("  Bibliotecas     : " + " · ".join("%s %s" % (dist, libs.get(dist) or "AUSENTE") for dist, _ in LIBS))
    out("  LibreOffice     : %s" % (soffice or "não encontrado (opcional: serve para converter em PDF e recalcular fórmulas)"))
    if missing:
        out("  FALTANDO        : %s" % ", ".join(missing))
        return 1
    out("  Tudo pronto. NÃO instale pacotes: rode seus scripts com `openclod-python script.py` e confira o resultado com")
    out("  `openclod-office inspect <arquivo>`.")
    return 0


# --------------------------------------------------------------------------------------- inspect: docx


def inspect_docx(path, full):
    import docx
    from docx.shared import Emu

    document = docx.Document(path)
    paragraphs = document.paragraphs
    info = {"kind": "docx", "headings": [], "tables": [], "warnings": []}
    for p in paragraphs:
        style = (p.style.name if p.style is not None else "") or ""
        if re.match(r"(?i)^(heading|título|titulo|title)", style) and p.text.strip():
            info["headings"].append({"style": style, "text": p.text.strip()})
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        info["tables"].append(
            {"rows": len(table.rows), "cols": len(table.columns), "header": rows[0] if rows else [], "preview": rows if full else rows[:6]}
        )
    words = sum(len(p.text.split()) for p in paragraphs) + sum(len(c.text.split()) for t in document.tables for r in t.rows for c in r.cells)
    info["paragraphs"] = len([p for p in paragraphs if p.text.strip()])
    info["words"] = words
    info["images"] = len(document.inline_shapes)
    section = document.sections[0]
    info["page"] = {
        "widthCm": round(Emu(section.page_width).cm, 1) if section.page_width else None,
        "heightCm": round(Emu(section.page_height).cm, 1) if section.page_height else None,
        "landscape": bool(section.page_width and section.page_height and section.page_width > section.page_height),
    }
    footer_xml = "".join(s.footer._element.xml for s in document.sections if s.footer is not None)
    header_text = " ".join(p.text for s in document.sections for p in s.header.paragraphs if p.text.strip())
    info["header"] = header_text[:120]
    info["footer"] = " ".join(p.text for s in document.sections for p in s.footer.paragraphs if p.text.strip())[:160]
    info["pageNumbers"] = bool(re.search(r"\bPAGE\b", footer_xml))
    run = longest = 0
    for p in paragraphs:
        run = run + 1 if not p.text.strip() and not p._element.xpath(".//w:drawing|.//w:br[@w:type='page']") else 0
        longest = max(longest, run)
    all_text = "\n".join(p.text for p in paragraphs)
    info["outline"] = [("  " if h["style"].endswith(("2", "3")) else "") + h["text"] for h in info["headings"]]
    if full:
        info["text"] = all_text
    warn = info["warnings"]
    if not info["headings"]:
        warn.append("nenhum título com estilo Heading/Título: o documento não tem estrutura (navegação, sumário).")
    if longest > 2:
        warn.append("%d parágrafos vazios seguidos: use espaçamento do parágrafo ou quebra de página em vez de Enter." % longest)
    if MARKDOWN_LEAK.search(all_text):
        warn.append("sobrou sintaxe de Markdown no texto (#, **, - [ ]): converta em estilos do Word.")
    if PLACEHOLDER_TEXT.search(all_text):
        warn.append("há texto de preenchimento (Lorem ipsum, TODO, XXX, [inserir…]).")
    if words < 40:
        warn.append("o documento tem só %d palavras." % words)
    return info


# --------------------------------------------------------------------------------------- inspect: xlsx


def inspect_xlsx(path, full):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    cached = openpyxl.load_workbook(path, data_only=True)
    info = {"kind": "xlsx", "sheets": [], "warnings": [], "charts": 0, "formulas": 0}
    with zipfile.ZipFile(path) as z:
        info["charts"] = len([n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n)])
    for ws in wb.worksheets:
        values = cached[ws.title]
        formulas = []
        no_cache = 0
        errors = []
        text_numbers = 0
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    shown = values[cell.coordinate].value
                    formulas.append({"ref": cell.coordinate, "formula": value, "cached": shown})
                    if shown is None:
                        no_cache += 1
                    elif isinstance(shown, str) and shown in ERROR_VALUES:
                        errors.append("%s!%s = %s" % (ws.title, cell.coordinate, shown))
                elif isinstance(value, str) and re.fullmatch(r"\s*-?[\d.]+(,\d+)?\s*", value) and re.search(r"\d", value):
                    text_numbers += 1
                elif isinstance(value, str) and value in ERROR_VALUES:
                    errors.append("%s!%s = %s" % (ws.title, cell.coordinate, value))
        sheet = {
            "name": ws.title,
            "dimensions": ws.dimensions,
            "rows": ws.max_row,
            "cols": ws.max_column,
            "formulas": len(formulas),
            "formulaSamples": formulas if full else formulas[:6],
            "withoutCachedValue": no_cache,
            "freezePanes": ws.freeze_panes,
            "autoFilter": ws.auto_filter.ref,
            "merged": len(ws.merged_cells.ranges),
            "conditionalFormats": len(ws.conditional_formatting),
            "tables": len(ws.tables),
        }
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1), [])]
        sheet["header"] = [h for h in header if h is not None][:12]
        info["sheets"].append(sheet)
        info["formulas"] += len(formulas)
        if errors:
            info["warnings"].append("valores de erro em %s: %s" % (ws.title, ", ".join(errors[:6])))
        if no_cache:
            info["warnings"].append(
                "%s: %d fórmula(s) sem valor em cache (visualizadores que não calculam mostram vazio; o Excel/LibreOffice calcula ao abrir). "
                "Com XlsxWriter passe value= em write_formula." % (ws.title, no_cache)
            )
        if text_numbers:
            info["warnings"].append("%s: %d célula(s) com número guardado como texto (grave int/float)." % (ws.title, text_numbers))
    names = [ws.title for ws in wb.worksheets]
    if len(names) > 1 and not any(re.search(r"(?i)resumo|summary", n) for n in names):
        info["notes"] = ["nenhuma aba chamada Resumo"]
    return info


# --------------------------------------------------------------------------------------- inspect: pptx


def _overflow_ratio(shape):
    """Rough estimate of text height over shape height (>1 = the text may spill out of the box)."""
    frame = shape.text_frame
    if not frame.text.strip() or not shape.width or not shape.height:
        return None
    width_pt = max(shape.width / EMU_PER_INCH * 72.0 - 14.0, 10.0)
    height_pt = max(shape.height / EMU_PER_INCH * 72.0 - 7.0, 10.0)
    total = 0.0
    for p in frame.paragraphs:
        size = None
        for r in p.runs:
            if r.font.size:
                size = r.font.size.pt
                break
        if size is None and p.font.size:
            size = p.font.size.pt
        size = size or 18.0
        per_line = max(1.0, width_pt / (size * 0.52))
        lines = max(1, math.ceil(len(p.text) / per_line)) if p.text else 1
        total += lines * size * 1.2 + (p.space_after.pt if p.space_after else 0)
    return total / height_pt


def inspect_pptx(path, full):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    prs = Presentation(path)
    info = {
        "kind": "pptx",
        "slideSize": {"widthCm": round(prs.slide_width / EMU_PER_CM, 1), "heightCm": round(prs.slide_height / EMU_PER_CM, 1)},
        "slides": [],
        "warnings": [],
        "charts": 0,
    }
    warn = info["warnings"]
    for number, slide in enumerate(prs.slides, start=1):
        item = {"n": number, "layout": slide.slide_layout.name, "title": None, "shapes": {"text": 0, "picture": 0, "table": 0, "chart": 0, "other": 0}, "text": []}
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            item["title"] = slide.shapes.title.text_frame.text.strip() or None
        for shape in slide.shapes:
            kind = "other"
            if shape.has_chart if hasattr(shape, "has_chart") else False:
                kind = "chart"
                info["charts"] += 1
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                kind = "picture"
            elif getattr(shape, "has_table", False) and shape.has_table:
                kind = "table"
            elif shape.has_text_frame and shape.text_frame.text.strip():
                kind = "text"
            item["shapes"][kind] += 1
            if shape.has_text_frame and shape.text_frame.text.strip():
                item["text"].append(shape.text_frame.text.strip().replace("\n", " / ")[:110 if not full else 600])
                ratio = _overflow_ratio(shape)
                if ratio and ratio > 1.12:
                    warn.append("slide %d: o texto de \"%s\" pode transbordar a caixa (~%d%% da altura); reduza o texto ou a fonte." % (number, shape.name, round(ratio * 100)))
            elif shape.is_placeholder and shape.placeholder_format.type not in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE):
                if kind == "other" and not getattr(shape, "has_table", False):
                    warn.append("slide %d: placeholder vazio \"%s\" (preencha ou apague)." % (number, shape.name))
            if shape.left is not None and shape.width is not None and (shape.left < -1000 or shape.left + shape.width > prs.slide_width + 1000):
                warn.append("slide %d: \"%s\" sai da área do slide na horizontal." % (number, shape.name))
            if shape.top is not None and shape.height is not None and (shape.top < -1000 or shape.top + shape.height > prs.slide_height + 1000):
                warn.append("slide %d: \"%s\" sai da área do slide na vertical." % (number, shape.name))
        item["notes"] = bool(slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip())
        if not item["title"]:
            warn.append("slide %d sem título (use o placeholder de título: navegação e acessibilidade)." % number)
        if PLACEHOLDER_TEXT.search(" ".join(item["text"])):
            warn.append("slide %d tem texto de preenchimento." % number)
        info["slides"].append(item)
    return info


# --------------------------------------------------------------------------------------- inspect: entry point


def render_inspection(info, size):
    lines = ["%s · %.0f KB" % (info["kind"].upper(), size / 1024.0)]
    if info["kind"] == "docx":
        lines.append("  %d parágrafos com texto, %d palavras, %d tabela(s), %d imagem(ns); página %s×%s cm%s" % (
            info["paragraphs"], info["words"], len(info["tables"]), info["images"], info["page"]["widthCm"], info["page"]["heightCm"],
            " (paisagem)" if info["page"]["landscape"] else ""))
        lines.append("  numeração de página: %s · cabeçalho: %s · rodapé: %s" % ("sim" if info["pageNumbers"] else "não", info["header"] or "—", info["footer"] or "—"))
        lines.append("  títulos:")
        lines.extend("    " + h for h in info["outline"][:30])
        for number, table in enumerate(info["tables"], start=1):
            lines.append("  tabela %d (%d×%d): %s" % (number, table["rows"], table["cols"], " | ".join(table["header"])))
            lines.extend("    " + " | ".join(row) for row in table["preview"][1:])
    elif info["kind"] == "xlsx":
        lines.append("  %d aba(s), %d fórmula(s), %d gráfico(s)" % (len(info["sheets"]), info["formulas"], info["charts"]))
        for sheet in info["sheets"]:
            lines.append("  aba \"%s\" %s: %d fórmula(s)%s%s" % (
                sheet["name"], sheet["dimensions"], sheet["formulas"],
                ", painéis congelados em %s" % sheet["freezePanes"] if sheet["freezePanes"] else "",
                ", filtro %s" % sheet["autoFilter"] if sheet["autoFilter"] else ""))
            if sheet["header"]:
                lines.append("    cabeçalho: " + " | ".join(str(h) for h in sheet["header"]))
            for f in sheet["formulaSamples"]:
                lines.append("    %s  %s  → %s" % (f["ref"], f["formula"], f["cached"] if f["cached"] is not None else "(sem valor em cache)"))
    else:
        lines.append("  %d slide(s), %d gráfico(s), %s×%s cm" % (len(info["slides"]), info["charts"], info["slideSize"]["widthCm"], info["slideSize"]["heightCm"]))
        for slide in info["slides"]:
            kinds = ", ".join("%d %s" % (n, k) for k, n in slide["shapes"].items() if n)
            lines.append("  %d. %s  [%s]%s" % (slide["n"], slide["title"] or "(sem título)", kinds, " · notas do apresentador" if slide["notes"] else ""))
    if info["warnings"]:
        lines.append("AVISOS:")
        lines.extend("  - " + w for w in info["warnings"])
    else:
        lines.append("Sem avisos.")
    return "\n".join(lines)


def cmd_inspect(args):
    path = os.path.abspath(args.file)
    if not os.path.isfile(path):
        out("arquivo não encontrado: " + path)
        return 2
    ext = os.path.splitext(path)[1].lower()
    handlers = {".docx": inspect_docx, ".xlsx": inspect_xlsx, ".pptx": inspect_pptx}
    if ext not in handlers:
        out("formato não suportado: %s (use .docx, .xlsx ou .pptx)" % ext)
        return 2
    if not zipfile.is_zipfile(path):
        out("ERRO: %s não é um arquivo %s válido (não é um zip)." % (os.path.basename(path), ext))
        return 1
    try:
        info = handlers[ext](path, args.full)
    except Exception as exc:
        out("ERRO: não consegui abrir %s como %s: %s: %s" % (os.path.basename(path), ext, type(exc).__name__, exc))
        return 1
    info["file"] = path
    info["bytes"] = os.path.getsize(path)
    info["ok"] = True
    if args.json:
        out(json.dumps(info, ensure_ascii=False, indent=2, default=str))
    else:
        out(render_inspection(info, info["bytes"]))
    return 0


# --------------------------------------------------------------------------------------- LibreOffice


_RECALC_XCU = (
    '<?xml version="1.0" encoding="UTF-8"?>\n'
    '<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">\n'
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>\n'
    "</oor:items>\n"
)


def soffice_convert(soffice, source, fmt, outdir, recalc=False):
    """Converts with LibreOffice in a throwaway profile. recalc=True makes it recalculate every formula on load
    (by default it trusts the values cached in an .xlsx, which would hide wrong or stale ones)."""
    profile = tempfile.mkdtemp(prefix="openclod-lo-")
    try:
        if recalc:
            os.makedirs(os.path.join(profile, "user"), exist_ok=True)
            with open(os.path.join(profile, "user", "registrymodifications.xcu"), "w", encoding="utf-8") as handle:
                handle.write(_RECALC_XCU)
        command = [soffice, "-env:UserInstallation=file:///" + profile.replace("\\", "/"), "--headless", "--convert-to", fmt, "--outdir", outdir, source]
        flags = 0x08000000 if os.name == "nt" else 0  # CREATE_NO_WINDOW
        subprocess.run(command, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=240, creationflags=flags)
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    produced = os.path.join(outdir, os.path.splitext(os.path.basename(source))[0] + "." + fmt.split(":")[0])
    return produced if os.path.isfile(produced) else None


def pdf_pages(path):
    with open(path, "rb") as handle:
        return len(re.findall(rb"/Type\s*/Page[^s]", handle.read()))


def cmd_pdf(args):
    soffice = find_soffice()
    if not soffice:
        out("LibreOffice não encontrado: não dá para converter em PDF aqui. Confira o arquivo com `openclod-office inspect`.")
        return 3
    source = os.path.abspath(args.file)
    outdir = os.path.abspath(args.out) if args.out else os.path.dirname(source)
    os.makedirs(outdir, exist_ok=True)
    produced = soffice_convert(soffice, source, "pdf", outdir)
    if not produced:
        out("ERRO: o LibreOffice não conseguiu abrir %s (arquivo corrompido?)." % os.path.basename(source))
        return 1
    out("PDF: %s · %d página(s)" % (produced, pdf_pages(produced)))
    return 0


def cmd_recalc(args):
    import openpyxl

    soffice = find_soffice()
    if not soffice:
        out("LibreOffice não encontrado: não dá para recalcular. Use `openclod-office inspect` para listar as fórmulas e confira-as.")
        return 3
    source = os.path.abspath(args.file)
    work = tempfile.mkdtemp(prefix="openclod-recalc-")
    try:
        produced = soffice_convert(soffice, source, "xlsx:Calc MS Excel 2007 XML", work, recalc=True)
        if not produced:
            out("ERRO: o LibreOffice não conseguiu abrir %s." % os.path.basename(source))
            return 1
        formulas = openpyxl.load_workbook(source)
        before = openpyxl.load_workbook(source, data_only=True)
        after = openpyxl.load_workbook(produced, data_only=True)
        total = errors = stale = 0
        lines = []
        for ws in formulas.worksheets:
            if ws.title not in after.sheetnames:
                continue
            for row in ws.iter_rows():
                for cell in row:
                    if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                        continue
                    total += 1
                    new = after[ws.title][cell.coordinate].value
                    old = before[ws.title][cell.coordinate].value
                    if new is None or (isinstance(new, str) and new in ERROR_VALUES):
                        errors += 1
                        lines.append("  ERRO  %s!%s  %s  → %s" % (ws.title, cell.coordinate, cell.value, new))
                    elif isinstance(old, (int, float)) and isinstance(new, (int, float)) and abs(old - new) > 1e-6 * max(1.0, abs(new)):
                        stale += 1
                        lines.append("  CACHE %s!%s  %s  em cache %s, recalculado %s" % (ws.title, cell.coordinate, cell.value, old, new))
        out("%d fórmula(s) recalculadas pelo LibreOffice: %d com erro, %d com valor em cache diferente." % (total, errors, stale))
        for line in lines[:40]:
            out(line)
        return 1 if errors else 0
    finally:
        shutil.rmtree(work, ignore_errors=True)


# --------------------------------------------------------------------------------------- main


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            pass
    parser = argparse.ArgumentParser(prog="openclod-office", description="Confere o ambiente e os documentos .docx/.xlsx/.pptx do Openclod.")
    sub = parser.add_subparsers(dest="command", required=True)
    p = sub.add_parser("check", help="bibliotecas disponíveis")
    p.add_argument("--json", action="store_true")
    p = sub.add_parser("inspect", help="o que há dentro de um .docx/.xlsx/.pptx")
    p.add_argument("file")
    p.add_argument("--json", action="store_true")
    p.add_argument("--full", action="store_true", help="todas as linhas/fórmulas/textos")
    p = sub.add_parser("pdf", help="converte para PDF (LibreOffice)")
    p.add_argument("file")
    p.add_argument("--out")
    p = sub.add_parser("recalc", help="recalcula as fórmulas de um .xlsx (LibreOffice)")
    p.add_argument("file")
    args = parser.parse_args(argv)
    return {"check": cmd_check, "inspect": cmd_inspect, "pdf": cmd_pdf, "recalc": cmd_recalc}[args.command](args)


if __name__ == "__main__":
    sys.exit(main())
